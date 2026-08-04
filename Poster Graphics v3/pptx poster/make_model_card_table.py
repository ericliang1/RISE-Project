"""Build a simple model-card table slide (original vs physics-guided parameter counts)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = "08_model_card_table.pptx"

HEADER = ["", "DeepSets", "GNN", "Set Transformer"]
ROWS = [
    ["Aggregation mechanism",
     "shared MLP + mean/max pool",
     "message passing, wind-aligned edges",
     "attention (2× ISAB + PMA)"],
    ["Original network parameters", "4,870,080", "5,610,048", "6,528,448"],
    ["Physics head parameters", "+10,433", "+10,433", "+10,433"],
    ["New network total", "4,880,513", "5,620,481", "6,538,881"],
    ["Parameter increase", "+0.21%", "+0.19%", "+0.16%"],
]
BOLD_ROWS = {4}  # "New network total" (1-indexed within body rows)

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark blue header
LIGHT = RGBColor(0xEB, 0xF1, 0xF7)    # banded row fill
TEXT = RGBColor(0x21, 0x21, 0x21)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# Title
title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
tf = title_box.text_frame
tf.text = "Model Parameter Comparison"
p = tf.paragraphs[0]
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = ACCENT
p.alignment = PP_ALIGN.CENTER

# Table
n_rows, n_cols = len(ROWS) + 1, len(HEADER)
tbl_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.8), Inches(1.6),
                                   Inches(11.7), Inches(4.6))
table = tbl_shape.table
table.columns[0].width = Inches(3.1)
for c in range(1, n_cols):
    table.columns[c].width = Inches(2.87)

def style_cell(cell, text, bold=False, header=False, fill=None):
    cell.text = text
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    for para in cell.text_frame.paragraphs:
        para.alignment = PP_ALIGN.CENTER
        for run in para.runs or [para.add_run()]:
            run.font.size = Pt(16)
            run.font.bold = bold or header
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if header else TEXT

for c, text in enumerate(HEADER):
    style_cell(table.cell(0, c), text, header=True, fill=ACCENT)

for r, row in enumerate(ROWS, start=1):
    band = LIGHT if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
    for c, text in enumerate(row):
        style_cell(table.cell(r, c), text, bold=(r in BOLD_ROWS or c == 0), fill=band)
        if c == 0:
            table.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

prs.save(OUT)
print(f"saved {OUT}")
